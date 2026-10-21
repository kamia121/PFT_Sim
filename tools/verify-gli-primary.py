#!/usr/bin/env python3
"""
Compares the GLI-2021 volume and GLI-2017 TLCO data used by PFT_Sim with the PRIMARY sources, as supplied in the
"GLI primary verification packet" (retrieved from the ERS publisher on 2026-10-20; see its provenance.json for URLs and hashes):

  1. The regression coefficients printed in Hall 2021 (Eur Respir J 57:2000289), Table 3, and in the 2020 correction to Stanojevic 2017
     (Eur Respir J 56:1750010), corrected Table 2. They are transcribed below from the printed tables as read by the packet's author
     (coefficient_facts.txt). The transcription is data in this file; it is not code from any package.
  2. The spline lookup tables published by the ERS with those papers (Hall2021_Lookup_tables.xlsx and the corrected Stanojevic workbooks).

Scope. This is a comparison of stored numbers with the printed coefficients and the published lookup tables. It is NOT a validation of
units, of interpolation between the 0.25-year rows, of behavior at the age boundaries, or of the LMS calculation (src/reference/gli.js)
against an official calculator. No official online-calculator case has been run (the calculator is behind a CAPTCHA and a paid key).
The printed tables were inspected visually in the publisher PDFs (Hall 2021 Table 3, PDF page 5; 2020 correction to Stanojevic 2017,
corrected Table 2, PDF page 2) and transcribed by the packet's author; the transcription is the weak link of this check. A separate Codex audit (2026-10-20)
matched 120/120 coefficient values and recomputed 100 measurements with the stored splines to 1e-12, but could not rerun the workbook
comparison (403), so the spline comparison is this script's alone. Source hashes are in src/reference/sources.json and in
docs/verification/gli-packet-2026-10-20/provenance.json.

Usage: python3 tools/verify-gli-primary.py PACKET_DIR
Exit status 1 if any coefficient differs or any spline value differs by more than 1e-9.
"""
import json, math, sys, pathlib
import openpyxl

PACKET = pathlib.Path(sys.argv[1])
DATA = pathlib.Path(__file__).resolve().parent.parent / 'src' / 'reference' / 'data'

# (const, ageLn, age, htLn, ht) of M; (const, ageLn, age) of S; L constant. Printed form in the comment.
HALL = {
  # measure: { sex: (M, S, L, sSpline) }
  'frc':    {'male': ((-13.4898, 0.1111, 0, 2.7634, 0),       (-1.60197, 0.01513, 0),         0.3416, True),
             'female': ((-12.7674, 0.1251, 0, 2.6049, 0),     (-1.48310, -0.03372, 0),        0.2898, True)},
  'tlc':    {'male': ((-10.5861, 0.1433, 0, 2.3155, 0),       (-2.0616143, 0, -0.0008534),    0.9337, True),
             'female': ((-10.1128, 0.1062, 0, 2.2259, 0),     (-2.0999321, 0, 0.0001564),     0.4636, True)},
  'rv':     {'male': ((-2.37211, 0, 0.01346, 0, 0.01307),     (-0.878572, 0, -0.007032),      0.5931, True),
             'female': ((-2.50593, 0, 0.01307, 0, 0.01379),   (-0.902550, 0, -0.006005),      0.4197, True)},
  'rv_tlc': {'male': ((2.634, 0, 0.01302, 0, -0.00008862),    (-0.96804, 0, -0.01004),        0.8646, True),
             'female': ((2.666, 0, 0.01411, 0, -0.00003689),  (-0.976602, 0, -0.009679),      0.8037, True)},
  'erv':    {'male': ((-17.328650, 0, -0.006288, 3.478116, 0), (-1.307616, 0, 0.009177),      0.5517, False),
             'female': ((-14.145513, 0, -0.009573, 2.871446, 0), (-1.54992, 0, 0.01409),      0.5326, False)},
  'ic':     {'male': ((-10.121688, 0, 0.001265, 2.188801, 0), (-1.856546, 0, 0.002008),       1.146, False),
             'female': ((-9.4438787, 0, -0.0002484, 2.0312769, 0), (-1.775276, 0, 0.002673),  0.9726, False)},
  'vc':     {'male': ((-10.134371, 0, -0.003532, 2.307980, 0), (-2.1367411, 0, 0.0009367),    0.8611, False),
             'female': ((-9.230600, 0, -0.005517, 2.116822, 0), (-2.220260, 0, 0.002956),     1.038, False)},
}
# Traditional units (mL/min/mmHg, mL/min/mmHg/L, L): M (const, ageLn, htLn), S (const, ageLn), L
STANOJEVIC = {
  'dlco': {'male': ((-7.034920, -0.012425, 2.018368), (-1.98996, 0.03536), 0.39482),
           'female': ((-5.159451, -0.015390, 1.618697), (-1.82905, -0.01815), 0.24160)},
  'kco':  {'male': ((4.088408, -0.113166, -0.415334), (-1.98186, 0.01460), 0.67330),
           'female': ((5.131492, -0.097395, -0.645656), (-1.63787, -0.07757), 0.48963)},
  'va':   {'male': ((-11.086573, 0.097047, 2.430021), (-2.20953, 0.01937), 0.62559),
           'female': ((-9.873970, 0.082868, 2.182316), (-2.08839, -0.01334), 0.51919)},
}
# SI intercepts printed alongside (TLCO, KCO): used only for the unit-consistency check
SI_INTERCEPT = {('dlco', 'male'): -8.129189, ('dlco', 'female'): -6.253720, ('kco', 'male'): 2.994137, ('kco', 'female'): 4.037222}

HALL_SHEET = {'frc': 'frc', 'tlc': 'tlc', 'rv': 'rv', 'rv_tlc': 'rvtlc', 'erv': 'erv', 'ic': 'ic', 'vc': 'vc'}
problems = []
lines = []

def close(a, b, tol):
    return abs(a - b) <= tol

def check_coef(label, stored, printed, tol=1e-12):
    for k, (s, p) in zip(('const', 'ageLn', 'age', 'htLn', 'ht'), zip(stored, printed)):
        if not close(s, p, tol):
            problems.append(f'{label}: {k} stored {s} printed {p}')

def vec(d):
    return (d['const'], d['ageLn'], d['age'], d['htLn'], d['ht'])

# 1. coefficients, volumes
vol = json.load(open(DATA / 'gli_volumes.json'))['measures']
for measure, by_sex in HALL.items():
    for sex, (m, s, l, has_s_spline) in by_sex.items():
        e = vol[measure][sex]
        check_coef(f'volumes {measure} {sex} M', vec(e['m']), m)
        check_coef(f'volumes {measure} {sex} S', vec(e['s']), (s[0], s[1], s[2], 0, 0))
        if not close(e['l']['const'], l, 1e-12) or e['l']['ageLn'] or e['l']['age']:
            problems.append(f'volumes {measure} {sex} L: stored {e["l"]} printed {l}')
        if bool(e['sSpline']) != has_s_spline:
            problems.append(f'volumes {measure} {sex}: sSpline flag stored {e["sSpline"]} expected {has_s_spline}')
        lines.append(f'coefficients ok: volumes {measure} {sex}')

# 1b. coefficients, TLCO (traditional units)
tl = json.load(open(DATA / 'gli_tlco.json'))['measures']
for measure, by_sex in STANOJEVIC.items():
    for sex, (m, s, l) in by_sex.items():
        e = tl[measure][sex]
        check_coef(f'tlco {measure} {sex} M', vec(e['m']), (m[0], m[1], 0, m[2], 0))
        check_coef(f'tlco {measure} {sex} S', vec(e['s']), (s[0], s[1], 0, 0, 0))
        if not close(e['l']['const'], l, 1e-12) or e['l']['ageLn'] or e['l']['age']:
            problems.append(f'tlco {measure} {sex} L: stored {e["l"]} printed {l}')
        lines.append(f'coefficients ok: tlco {measure} {sex}')

# 1c. unit consistency: traditional intercept minus SI intercept is the unit-conversion constant
for (measure, sex), si in SI_INTERCEPT.items():
    trad = STANOJEVIC[measure][sex][0][0]
    lines.append(f'unit offset {measure} {sex}: traditional - SI intercept = {trad - si:.6f} (factor {math.exp(trad - si):.5f})')

# 2. splines against the published workbooks
def compare_sheet(label, ws, spl, has_s, first_col):
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    rows = [r[first_col:first_col + 4] for r in rows if r[first_col] is not None]
    if len(rows) != len(spl['age']):
        problems.append(f'{label}: workbook has {len(rows)} ages, stored {len(spl["age"])}')
    worst = {'age': 0.0, 'm': 0.0, 's': 0.0, 'l': 0.0}
    for r, a, m, s, l in zip(rows, spl['age'], spl['m'], spl['s'], spl['l']):
        worst['age'] = max(worst['age'], abs(r[0] - a))
        worst['m'] = max(worst['m'], abs(r[1] - m))
        worst['l'] = max(worst['l'], abs((r[3] or 0) - l))
        if has_s:
            worst['s'] = max(worst['s'], abs((r[2] or 0) - s))
    for k, v in worst.items():
        if v > 1e-9:
            problems.append(f'{label}: max |{k}| difference {v:.3g}')
    lines.append(f'splines {label}: {len(rows)} ages, ages {rows[0][0]} to {rows[-1][0]}, max abs diff M {worst["m"]:.2g}, S {worst["s"]:.2g}, L {worst["l"]:.2g}')

wb = openpyxl.load_workbook(PACKET / 'Hall2021_Lookup_tables.xlsx', data_only=True)
for measure, by_sex in HALL.items():
    for sex, (_, _, _, has_s) in by_sex.items():
        sheet = f'{HALL_SHEET[measure]}_{"m" if sex == "male" else "f"}_lookuptable'
        compare_sheet(f'volumes {measure} {sex}', wb[sheet], vol[measure][sex]['splines'], has_s, 0)

wb = openpyxl.load_workbook(PACKET / 'Stanojevic2017_corrected_Lookup_Traditional.xlsx', data_only=True)
for measure in STANOJEVIC:
    for sex in ('male', 'female'):
        sheet = f'{measure.upper()}_{"m" if sex == "male" else "f"}'
        compare_sheet(f'tlco {measure} {sex}', wb[sheet], tl[measure][sex]['splines'], True, 1)

print('\n'.join(lines))
print()
if problems:
    print(f'{len(problems)} PROBLEM(S):')
    print('\n'.join(problems))
    sys.exit(1)
print('All coefficients equal the transcribed printed values (to 1e-12) and all spline tables equal the published workbooks (to 1e-9).')
print('Not checked here: units, interpolation, age boundaries, LMS evaluation against an official calculator (none was run).')
